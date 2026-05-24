import io
import logging
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

logger = logging.getLogger(__name__)


def export_submissions_xlsx(submissions: list[dict]) -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Сданные РТ"

    headers = ["№", "ФИО ученика", "Username", "Telegram ID", "Группа",
               "Куратор", "Дата сдачи", "Время сдачи", "Статус", "PDF файл"]

    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(bold=True, color="FFFFFF")

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for idx, s in enumerate(submissions, 1):
        dt_str = s.get("submitted_at", "")
        date_part, time_part = "", ""
        if dt_str:
            try:
                dt = datetime.fromisoformat(dt_str)
                date_part = dt.strftime("%d.%m.%Y")
                time_part = dt.strftime("%H:%M")
            except Exception:
                date_part = dt_str

        pdf_name = ""
        if s.get("pdf_path"):
            import os
            pdf_name = os.path.basename(s["pdf_path"])

        ws.append([
            idx,
            s.get("student_full_name", ""),
            f"@{s.get('username', '')}" if s.get("username") else "",
            s.get("student_telegram_id", ""),
            s.get("group_title", ""),
            s.get("curator_name", ""),
            date_part,
            time_part,
            s.get("status", ""),
            pdf_name,
        ])

    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_bookings_xlsx(bookings: list[dict]) -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Записи на зачёт"

    headers = ["№", "ФИО ученика", "Username", "Telegram ID", "Группа",
               "Куратор", "Дата зачёта", "Время зачёта", "Google Meet", "Статус"]

    header_fill = PatternFill("solid", fgColor="70AD47")
    header_font = Font(bold=True, color="FFFFFF")

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for idx, b in enumerate(bookings, 1):
        ws.append([
            idx,
            b.get("student_full_name", ""),
            f"@{b.get('student_username', '')}" if b.get("student_username") else "",
            b.get("student_telegram_id", ""),
            b.get("group_title", ""),
            b.get("curator_name", ""),
            b.get("date", ""),
            b.get("slot_time", ""),
            b.get("google_meet_link", ""),
            b.get("status", ""),
        ])

    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
