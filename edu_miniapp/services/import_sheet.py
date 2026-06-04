"""Импорт журнала из таблицы (Google Sheets → выгрузка в .xlsx или .csv).

Ожидаемый формат (как обычно ведут журнал):

    A1          B1            C1           ...   <- строка 1: темы уроков
    (ФИО)       Клетка        ДНК          ...
    A2          B2            C2           ...   <- строка 2: максимальные баллы (числа)
    (пусто)     10            5            ...
    A3..        B3..          C3..         ...   <- строки учеников
    Иванов И.   8             5
    Петров П.   6             4

Логика распознавания:
  • строка 1 — темы уроков (со 2-го столбца);
  • если во 2-й строке (со 2-го столбца) стоят числа — это максимальные баллы;
    иначе максимум = DEFAULT_MAX, а 2-я строка считается учеником;
  • далее идут ученики: 1-й столбец — ФИО, остальные — баллы.

Как выгрузить из Google Sheets: Файл → Скачать → Microsoft Excel (.xlsx)
или CSV. Полученный файл отправьте боту.
"""
import csv
from pathlib import Path
from typing import Any

DEFAULT_MAX = 100.0


def _to_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    s = str(value).strip().replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _looks_like_max_row(cells: list[Any]) -> bool:
    """Вся строка (со 2-го столбца) состоит из чисел → это строка максимумов."""
    vals = cells[1:]
    nums = [_to_float(v) for v in vals]
    non_empty = [n for n, v in zip(nums, vals) if str(v).strip() != ""]
    return bool(non_empty) and all(n is not None for n in non_empty)


def _read_rows(path: str) -> list[list[Any]]:
    p = Path(path)
    ext = p.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook  # ленивый импорт
        wb = load_workbook(p, read_only=True, data_only=True)
        ws = wb.active
        rows = [[c for c in row] for row in ws.iter_rows(values_only=True)]
        wb.close()
        return rows
    if ext in (".csv", ".tsv"):
        delimiter = "\t" if ext == ".tsv" else ","
        with open(p, encoding="utf-8-sig", newline="") as f:
            return [row for row in csv.reader(f, delimiter=delimiter)]
    raise ValueError(f"Неподдерживаемый формат: {ext}. Нужен .xlsx или .csv")


def parse_sheet(path: str) -> dict[str, Any]:
    """Разобрать файл в структуру {lessons:[...], students:[...]}."""
    rows = [r for r in _read_rows(path) if r and any(str(c).strip() for c in r)]
    if len(rows) < 2:
        raise ValueError("В таблице слишком мало данных")

    header = rows[0]
    topics = [str(c).strip() for c in header[1:] if str(c).strip() != ""]
    n = len(topics)
    if n == 0:
        raise ValueError("Не найдены темы уроков в первой строке")

    idx = 1
    max_scores = [DEFAULT_MAX] * n
    if idx < len(rows) and _looks_like_max_row(rows[idx]):
        raw = [_to_float(c) for c in rows[idx][1 : 1 + n]]
        max_scores = [(m if m is not None else DEFAULT_MAX) for m in raw]
        idx += 1

    students = []
    for row in rows[idx:]:
        name = str(row[0]).strip() if row and row[0] is not None else ""
        if not name:
            continue
        scores = [_to_float(c) for c in row[1 : 1 + n]]
        scores += [None] * (n - len(scores))  # добить до n
        students.append({"name": name, "scores": scores})

    lessons = [{"topic": t, "max_score": m} for t, m in zip(topics, max_scores)]
    return {"lessons": lessons, "students": students}


def _norm(name: str) -> str:
    return " ".join(name.lower().split())
